"""
Resume parser — supports PDF (via PyMuPDF) and DOCX (via python-docx).
Extracts: contact info, education, skills, projects, experience,
          certifications, achievements, and summary.
"""
import re
import io
from typing import Optional

try:
    import fitz  # PyMuPDF
    HAS_FITZ = True
except ImportError:
    HAS_FITZ = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ─── Section header keywords ──────────────────────────────────────────────────

SECTION_HEADERS = {
    "contact": ["contact", "personal information", "personal details", "about me"],
    "education": ["education", "academic background", "qualifications", "academic history"],
    "skills": [
        "skills", "technical skills", "technologies", "tech stack",
        "competencies", "areas of expertise", "core competencies",
    ],
    "experience": [
        "experience", "work experience", "employment history", "professional experience",
        "internships", "internship experience", "work history",
    ],
    "projects": ["projects", "personal projects", "academic projects", "key projects", "project work"],
    "certifications": ["certifications", "certificates", "courses", "training", "licenses"],
    "achievements": [
        "achievements", "awards", "honors", "accomplishments",
        "extra-curricular", "extracurricular", "activities",
    ],
    "summary": ["summary", "profile", "objective", "about", "professional summary", "career objective"],
}

# ─── Regexes ──────────────────────────────────────────────────────────────────

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s\-.]?)?(?:\(?\d{2,4}\)?[\s\-.]?)?\d{3,4}[\s\-.]?\d{4}")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w\-]+/?", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[\w\-]+(?:/[\w\-]+)?/?", re.I)
URL_RE = re.compile(r"https?://[^\s]+", re.I)
CGPA_RE = re.compile(r"\b(?:cgpa|gpa|grade)\s*:?\s*(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)", re.I)
DATE_RE = re.compile(
    r"\b(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
    r"jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
    r"\s*\.?\s*\d{4}\b"
    r"|\b\d{4}\s*[-–]\s*(?:\d{4}|present|current|ongoing)\b"
    r"|\b\d{4}\b",
    re.I,
)
BULLET_RE = re.compile(r"^[\s]*[•\-\*\u2022\u2023\u25e6\u2043\u2219▪▸►]\s+", re.M)


# ─── Text extraction ──────────────────────────────────────────────────────────

def extract_text_from_pdf(file_bytes: bytes) -> str:
    if not HAS_FITZ:
        raise ImportError("PyMuPDF (fitz) is not installed")
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages = []
    for page in doc:
        pages.append(page.get_text("text"))
    doc.close()
    return "\n".join(pages)


def extract_text_from_docx(file_bytes: bytes) -> str:
    if not HAS_DOCX:
        raise ImportError("python-docx is not installed")
    doc = DocxDocument(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also grab table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append(text)
    return "\n".join(paragraphs)


def extract_text(file_bytes: bytes, filename: str) -> str:
    fname = filename.lower()
    if fname.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif fname.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}")


# ─── Section detection ────────────────────────────────────────────────────────

def detect_section(line: str) -> Optional[str]:
    """Return section key if the line looks like a section header, else None."""
    clean = line.strip().lower().rstrip(":")
    # Short lines that match a known header are considered headers
    if len(clean) > 60:
        return None
    for key, patterns in SECTION_HEADERS.items():
        for p in patterns:
            if clean == p or clean.startswith(p + " ") or clean.endswith(" " + p):
                return key
    return None


def split_into_sections(text: str) -> dict:
    """Split raw text into named sections based on header detection."""
    lines = text.splitlines()
    sections: dict = {"_header": []}
    current = "_header"

    for line in lines:
        section = detect_section(line)
        if section:
            current = section
            if current not in sections:
                sections[current] = []
        else:
            if current not in sections:
                sections[current] = []
            sections[current].append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items()}


# ─── Field extractors ─────────────────────────────────────────────────────────

def extract_contact(header_text: str, full_text: str) -> dict:
    """Extract contact fields from header section or full text."""
    search_text = (header_text + "\n" + full_text[:500]) if header_text else full_text[:600]

    email = EMAIL_RE.search(search_text)
    phone = PHONE_RE.search(search_text)
    linkedin = LINKEDIN_RE.search(search_text)
    github = GITHUB_RE.search(search_text)

    # Name heuristic: first non-empty line that is not a URL/email and looks like a name
    name = None
    for line in full_text.splitlines()[:8]:
        line = line.strip()
        if (
            line
            and not EMAIL_RE.search(line)
            and not PHONE_RE.search(line)
            and not URL_RE.search(line)
            and len(line.split()) <= 6
            and len(line) >= 3
        ):
            name = line
            break

    # Location: look for city/state pattern
    location = None
    loc_patterns = [
        re.compile(r"\b([A-Z][a-z]+(?:\s[A-Z][a-z]+)*,\s*[A-Z]{2,})\b"),
        re.compile(r"\b([A-Z][a-z]+(?:\s[A-Z][a-z]+)*,\s*(?:India|USA|UK|Canada|Australia))\b", re.I),
    ]
    for pat in loc_patterns:
        m = pat.search(search_text)
        if m:
            location = m.group(1)
            break

    # Portfolio: any URL that is not linkedin/github
    portfolio = None
    for url in URL_RE.findall(search_text):
        if "linkedin" not in url.lower() and "github" not in url.lower():
            portfolio = url
            break

    return {
        "name": name,
        "email": email.group() if email else None,
        "phone": phone.group().strip() if phone else None,
        "location": location,
        "linkedin": linkedin.group() if linkedin else None,
        "github": github.group() if github else None,
        "portfolio": portfolio,
    }


def clean_bullet_lines(text: str) -> list:
    """Split text into cleaned bullet-point lines or regular non-empty lines."""
    lines = []
    for line in text.splitlines():
        line = BULLET_RE.sub("", line).strip()
        if line:
            lines.append(line)
    return lines


def extract_education(section_text: str) -> list:
    """Parse education section into a list of education items."""
    if not section_text:
        return []
    items = []
    blocks = re.split(r"\n{2,}", section_text)
    for block in blocks:
        if not block.strip():
            continue
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        if not lines:
            continue
        item: dict = {}
        # First line is usually institution or degree
        item["institution"] = lines[0] if lines else None
        # Second line might be degree/field
        if len(lines) > 1:
            item["degree"] = lines[1]
        # Search for CGPA
        cgpa_m = CGPA_RE.search(block)
        if cgpa_m:
            item["cgpa"] = f"{cgpa_m.group(1)}/{cgpa_m.group(2)}"
        # Search for dates
        dates = DATE_RE.findall(block)
        if len(dates) >= 2:
            item["start_date"] = dates[0]
            item["end_date"] = dates[1]
        elif len(dates) == 1:
            item["end_date"] = dates[0]
        # Look for field of study (B.Tech, B.E., M.Tech, etc.)
        fos_match = re.search(
            r"\b(B\.?Tech|B\.?E|M\.?Tech|M\.?Sc|B\.?Sc|MBA|BCA|MCA|"
            r"Bachelor|Master|PhD|Diploma)[^,\n]*",
            block, re.I
        )
        if fos_match:
            item["field_of_study"] = fos_match.group().strip()
        # Coursework
        cw_match = re.search(r"(?:coursework|courses)\s*:?\s*(.+)", block, re.I)
        if cw_match:
            item["coursework"] = [c.strip() for c in re.split(r"[,;]", cw_match.group(1)) if c.strip()]
        if item.get("institution") or item.get("degree"):
            items.append(item)
    return items


def extract_skills(section_text: str) -> dict:
    """Parse skills into categorized lists."""
    if not section_text:
        return {}

    SKILL_CATEGORIES = {
        "programming": [
            "python", "java", "c++", "c", "javascript", "typescript", "go", "rust",
            "kotlin", "swift", "ruby", "php", "scala", "r", "matlab",
        ],
        "web_dev": [
            "react", "vue", "angular", "html", "css", "node", "express", "django",
            "flask", "fastapi", "spring", "next.js", "nuxt", "svelte", "tailwind",
            "bootstrap", "jquery", "graphql", "rest", "api",
        ],
        "database": [
            "mysql", "postgresql", "sqlite", "mongodb", "redis", "cassandra",
            "dynamodb", "oracle", "sql server", "mariadb", "firebase",
        ],
        "ai_ml": [
            "machine learning", "deep learning", "tensorflow", "pytorch", "keras",
            "scikit-learn", "pandas", "numpy", "nlp", "computer vision",
            "neural network", "llm", "transformers", "huggingface", "langchain",
        ],
        "data_analytics": [
            "tableau", "power bi", "excel", "data analysis", "data visualization",
            "matplotlib", "seaborn", "plotly", "spark", "hadoop",
        ],
        "cloud": [
            "aws", "azure", "gcp", "google cloud", "heroku", "vercel", "netlify",
            "docker", "kubernetes", "ci/cd", "terraform",
        ],
        "tools": [
            "git", "github", "gitlab", "jira", "linux", "bash", "postman",
            "vscode", "intellij", "figma", "notion", "confluence",
        ],
    }

    # Collect raw skill tokens
    raw_skills = []
    for line in section_text.splitlines():
        # split on common delimiters
        parts = re.split(r"[,;|·•\t]+", line)
        for p in parts:
            s = p.strip().strip(":").strip()
            if s and len(s) < 60:
                raw_skills.append(s)

    categorized: dict = {k: [] for k in SKILL_CATEGORIES}
    categorized["other"] = []

    for skill in raw_skills:
        assigned = False
        skill_lower = skill.lower()
        for cat, keywords in SKILL_CATEGORIES.items():
            if any(kw in skill_lower for kw in keywords):
                if skill not in categorized[cat]:
                    categorized[cat].append(skill)
                assigned = True
                break
        if not assigned and skill:
            if skill not in categorized["other"]:
                categorized["other"].append(skill)

    # Remove empty categories
    return {k: v for k, v in categorized.items() if v}


def extract_projects(section_text: str) -> list:
    """Parse projects section into list of project items."""
    if not section_text:
        return []
    items = []
    blocks = re.split(r"\n{2,}", section_text)
    for block in blocks:
        if not block.strip():
            continue
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        if not lines:
            continue
        item: dict = {"name": lines[0]}
        remaining = lines[1:]
        description_lines = []
        for line in remaining:
            # Tech stack line
            tech_m = re.search(r"(?:tech(?:nologies)?|stack|built with|tools?)\s*:?\s*(.+)", line, re.I)
            if tech_m:
                item["tech_stack"] = [t.strip() for t in re.split(r"[,|]", tech_m.group(1)) if t.strip()]
                continue
            # URLs
            if GITHUB_RE.search(line):
                item["github_url"] = GITHUB_RE.search(line).group()
                continue
            if URL_RE.search(line) and "github" not in line.lower():
                item["live_url"] = URL_RE.search(line).group()
                continue
            description_lines.append(line)
        item["description"] = " ".join(description_lines) if description_lines else None
        items.append(item)
    return items


def extract_experience(section_text: str) -> list:
    """Parse experience section."""
    if not section_text:
        return []
    items = []
    blocks = re.split(r"\n{2,}", section_text)
    for block in blocks:
        if not block.strip():
            continue
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        if not lines:
            continue
        item: dict = {}
        # First line is usually role @ company or company name
        first = lines[0]
        at_split = re.split(r"\s+(?:at|@|,)\s+|\s*[-–]\s*", first, maxsplit=1)
        if len(at_split) == 2:
            item["role"] = at_split[0].strip()
            item["company"] = at_split[1].strip()
        else:
            item["company"] = first
        # Dates
        dates = DATE_RE.findall(block)
        if len(dates) >= 2:
            item["start_date"] = dates[0]
            item["end_date"] = dates[1]
        elif len(dates) == 1:
            item["end_date"] = dates[0]
        # Responsibilities = bullet lines after the first
        responsibilities = clean_bullet_lines("\n".join(lines[1:]))
        if responsibilities:
            item["responsibilities"] = responsibilities
        if item.get("company") or item.get("role"):
            items.append(item)
    return items


def extract_certifications(section_text: str) -> list:
    """Parse certifications section."""
    if not section_text:
        return []
    items = []
    for line in section_text.splitlines():
        line = line.strip()
        line = BULLET_RE.sub("", line).strip()
        if not line:
            continue
        item: dict = {"name": line}
        url_m = URL_RE.search(line)
        if url_m:
            item["url"] = url_m.group()
            item["name"] = line[:url_m.start()].strip().rstrip("-–|:")
        date_m = DATE_RE.search(line)
        if date_m:
            item["date"] = date_m.group()
        # Issuer: "by <Issuer>" or "from <Issuer>"
        issuer_m = re.search(r"(?:by|from|issued by)\s+(.+?)(?:\s*[\|\-–]|\s*$)", line, re.I)
        if issuer_m:
            item["issuer"] = issuer_m.group(1).strip()
        if item.get("name"):
            items.append(item)
    return items


def extract_achievements(section_text: str) -> list:
    """Parse achievements section into list of strings."""
    if not section_text:
        return []
    return clean_bullet_lines(section_text)


def extract_summary(section_text: str) -> Optional[str]:
    if not section_text:
        return None
    return section_text.strip()


# ─── Main parse function ──────────────────────────────────────────────────────

def parse_resume(file_bytes: bytes, filename: str) -> dict:
    """
    Parse a resume file (PDF or DOCX) and return structured data.

    Returns a dict matching ParsedResumeData schema:
    {
        contact, education, skills, projects, experience,
        certifications, achievements, summary
    }
    """
    raw_text = extract_text(file_bytes, filename)
    sections = split_into_sections(raw_text)

    header_text = sections.get("_header", "")

    contact = extract_contact(header_text, raw_text)
    education = extract_education(sections.get("education", ""))
    skills = extract_skills(sections.get("skills", ""))
    projects = extract_projects(sections.get("projects", ""))
    experience = extract_experience(sections.get("experience", ""))
    certifications = extract_certifications(sections.get("certifications", ""))
    achievements = extract_achievements(sections.get("achievements", ""))
    summary = extract_summary(sections.get("summary", ""))

    return {
        "contact": contact,
        "education": education,
        "skills": skills,
        "projects": projects,
        "experience": experience,
        "certifications": certifications,
        "achievements": achievements,
        "summary": summary,
        "_raw_text": raw_text,
    }
